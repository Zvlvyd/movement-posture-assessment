"""将 第一次检查-技术文档.md 转换为 Word (.docx) 格式"""
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_shading(cell, color):
    """设置单元格背景色"""
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def add_code_block(doc, code_text):
    """添加代码块（灰色背景段落）"""
    for line in code_text.strip().split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        # 设置段落背景色
        shading = OxmlElement('w:shd')
        shading.set(qn('w:fill'), 'F5F5F5')
        shading.set(qn('w:val'), 'clear')
        p._element.get_or_add_pPr().append(shading)

def add_para(doc, text, bold=False, size=10.5, color=None, font_name=None, alignment=None, space_after=6, space_before=0):
    """添加段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    if font_name:
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    return p

def add_heading_styled(doc, text, level):
    """添加标题"""
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = '微软雅黑'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    return h

def convert_md_to_docx(md_path, docx_path):
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '微软雅黑'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # 页面设置
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.18)
    section.right_margin = Cm(3.18)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    in_code_block = False
    code_buffer = []
    in_table = False
    table_rows = []

    while i < len(lines):
        line = lines[i].rstrip()

        # 代码块
        if line.startswith('```'):
            if in_code_block:
                add_code_block(doc, '\n'.join(code_buffer))
                code_buffer = []
                in_code_block = False
            else:
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        # 空行
        if not line:
            if in_table and table_rows:
                # 结束表格并渲染
                render_table(doc, table_rows)
                table_rows = []
                in_table = False
            i += 1
            continue

        # 表格行
        if line.startswith('|') and line.endswith('|'):
            # 跳过分隔行
            if re.match(r'^\|[\s\-:|]+\|$', line):
                i += 1
                continue
            in_table = True
            cells = [c.strip() for c in line.split('|')[1:-1]]
            table_rows.append(cells)
            i += 1
            continue
        elif in_table and table_rows:
            # 非表格行结束表格
            render_table(doc, table_rows)
            table_rows = []
            in_table = False

        # 标题
        if line.startswith('# '):
            add_heading_styled(doc, line[2:], 0)
        elif line.startswith('## '):
            add_heading_styled(doc, line[3:], 1)
        elif line.startswith('### '):
            add_heading_styled(doc, line[4:], 2)
        elif line.startswith('#### '):
            add_heading_styled(doc, line[5:], 3)

        # 分隔线
        elif line == '---':
            doc.add_paragraph('─' * 60)

        # 无序列表
        elif line.startswith('- ') or line.startswith('  - '):
            text = line.lstrip('- ').strip()
            # 处理内联代码 `code`
            text = re.sub(r'`([^`]+)`', r'\1', text)
            # 处理粗体 **text**
            text = re.sub(r'\*\*([^*]+)\*\*', r'\1', text)
            p = doc.add_paragraph(style='List Bullet')
            p.clear()
            run = p.add_run(text)
            run.font.size = Pt(10.5)
            run.font.name = '微软雅黑'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

        # 普通段落（处理内联格式）
        else:
            # 处理链接 [text](url)
            line = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', line)
            # 处理粗体 **text**
            parts = re.split(r'(\*\*[^*]+\*\*)', line)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                    run.font.size = Pt(10.5)
                    run.font.name = '微软雅黑'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
                else:
                    # 处理行内代码
                    sub_parts = re.split(r'(`[^`]+`)', part)
                    for sp in sub_parts:
                        if sp.startswith('`') and sp.endswith('`'):
                            run = p.add_run(sp[1:-1])
                            run.font.name = 'Consolas'
                            run.font.size = Pt(9)
                        else:
                            run = p.add_run(sp)
                            run.font.size = Pt(10.5)
                            run.font.name = '微软雅黑'
                            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

        i += 1

    # 处理末尾未关闭的表格
    if in_table and table_rows:
        render_table(doc, table_rows)

    doc.save(docx_path)
    print(f"Word 文档已生成: {docx_path}")


def render_table(doc, rows):
    """渲染表格到文档"""
    if not rows:
        return

    num_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=num_cols, style='Table Grid')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for r_idx, row in enumerate(rows):
        for c_idx, cell_text in enumerate(row):
            if c_idx < num_cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = ''
                p = cell.paragraphs[0]
                run = p.add_run(cell_text)
                run.font.size = Pt(9)
                run.font.name = '微软雅黑'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

                # 表头行（第一行）加粗并灰色背景
                if r_idx == 0:
                    run.bold = True
                    run.font.size = Pt(9.5)
                    set_cell_shading(cell, 'E8E8E8')

    doc.add_paragraph()  # 表后空行


if __name__ == '__main__':
    convert_md_to_docx(
        r'd:\workbench\program3\docs\第一次检查-技术文档.md',
        r'd:\workbench\program3\docs\第一次检查-技术文档.docx'
    )
