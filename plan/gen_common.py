from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

today_str = datetime.date.today().strftime("%Y-%m-%d")

def make_doc():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.0); s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)
    return doc

def set_heading_style(doc):
    for lv in [1,2,3,4]:
        h = doc.styles["Heading " + str(lv)]
        if lv == 1: h.font.size = Pt(18); h.font.color.rgb = RGBColor(0x1A,0x56,0xDB)
        elif lv == 2: h.font.size = Pt(14); h.font.color.rgb = RGBColor(0x2C,0x3E,0x50)
        elif lv == 3: h.font.size = Pt(12); h.font.color.rgb = RGBColor(0x34,0x49,0x5E)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1+len(rows), cols=len(headers))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,th in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = th
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.bold = True; r.font.size = Pt(9)
        s = OxmlElement("w:shd"); s.set(qn("w:fill"),"1A56DB"); s.set(qn("w:val"),"clear")
        c._tc.get_or_add_tcPr().append(s)
        for p in c.paragraphs:
            for r in p.runs: r.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
    for ri,row in enumerate(rows):
        for ci,val in enumerate(row):
            c = t.rows[ri+1].cells[ci]; c.text = str(val)
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(9)
            if ri % 2 == 1:
                s = OxmlElement("w:shd"); s.set(qn("w:fill"),"EBF0FA"); s.set(qn("w:val"),"clear")
                c._tc.get_or_add_tcPr().append(s)
    doc.add_paragraph()
    return t

def add_code(doc, text):
    for line in text.strip().split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
        run = p.add_run(line); run.font.name = "Consolas"; run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x33,0x33,0x33)

def add_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text); run.font.size = Pt(9); run.font.color.rgb = RGBColor(0xE6,0x7E,0x22)
    run.italic = True

def cover(doc, title, subtitle, desc):
    for _ in range(5): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('运动姿态评估与纠错系统'); r.font.size = Pt(22); r.font.color.rgb = RGBColor(0x1A,0x56,0xDB)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.font.size = Pt(18); r.font.bold = True; r.font.color.rgb = RGBColor(0x2C,0x3E,0x50)
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle); r.font.size = Pt(14); r.font.color.rgb = RGBColor(0x34,0x49,0x5E)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(desc + '    ' + today_str); r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x95,0xA5,0xA6)
    doc.add_page_break()

def footer(doc):
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('--- 文档结束 ---'); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x95,0xA5,0xA6); r.italic = True

print("Common ready")